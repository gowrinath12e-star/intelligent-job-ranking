"""Downloads are genuine, correct and purpose-limited CSV/OOXML files."""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.ranking import rank_applicants
from utils.exports import breakdown_frame, csv_bytes, excel_bytes, ranked_frame, spreadsheet_safe, validation_workbook


def test_csv_matches_ranked_records(applicant_frame):
    result = rank_applicants(applicant_frame)
    frame = ranked_frame(result.ranked_candidates)
    exported = pd.read_csv(io.BytesIO(csv_bytes(frame)))
    assert exported["Candidate ID"].tolist() == [candidate["candidate_id"] for candidate in result.ranked_candidates]
    assert "email" not in [column.lower() for column in exported.columns]
    assert len(exported) == 2


def test_excel_sheets_and_order(applicant_frame):
    result = rank_applicants(applicant_frame)
    raw = excel_bytes({"Ranking": ranked_frame(result.ranked_candidates), "Shortlist": ranked_frame(result.shortlist)})
    workbook = load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames == ["Ranking", "Shortlist"]
    assert workbook["Ranking"]["B2"].value == result.ranked_candidates[0]["candidate_id"]
    assert workbook["Ranking"].freeze_panes == "A2"


def test_breakdown_has_four_rows_per_candidate(applicant_frame):
    result = rank_applicants(applicant_frame)
    frame = breakdown_frame(result.ranked_candidates)
    assert len(frame) == 8
    assert frame.groupby("Candidate ID")["Contribution (points)"].sum().tolist() == pytest.approx([candidate["final_score"] for candidate in result.ranked_candidates])


def test_validation_workbook(applicant_frame):
    applicant_frame.loc[1, "experience_years"] = -1.0
    result = rank_applicants(applicant_frame)
    workbook = load_workbook(io.BytesIO(validation_workbook(result.validation)))
    assert workbook.sheetnames == ["Summary", "Issues", "Duplicate IDs", "Cleaning"]
    assert workbook["Issues"].max_row >= 2


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1)", "-cmd", "@formula", "\t=1", "\r=1", "   =1"])
def test_spreadsheet_formula_escape(value):
    frame = spreadsheet_safe(pd.DataFrame({"label": [value], "number": [-3.5]}))
    assert frame.label.iloc[0] == "'" + value
    assert frame.number.iloc[0] == -3.5
    raw = excel_bytes({"Safe": pd.DataFrame({"label": [value]})})
    cell = load_workbook(io.BytesIO(raw))["Safe"]["A2"]
    assert cell.data_type != "f"


def test_empty_report_with_headings():
    assert len(csv_bytes(ranked_frame([]))) > 0
    workbook = load_workbook(io.BytesIO(excel_bytes({"Ranking": ranked_frame([])})))
    assert workbook["Ranking"].max_row == 1


def test_invalid_sheet_names():
    with pytest.raises(ValueError):
        excel_bytes({})
    with pytest.raises(ValueError):
        excel_bytes({"Bad/Name": pd.DataFrame()})
